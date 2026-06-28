from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables" / "team_presentations"
DOCX_PATH = OUT_DIR / "Tamaffo_Fabrice_Database_Project_Lead_Summary.docx"

GREEN_DARK = "123D32"
GREEN = "13795B"
GREEN_LIGHT = "DDF2E9"
MINT = "F0F8F4"
GOLD = "D59A16"
INK = "19352E"
TEXT = "334E47"
MUTED = "60756E"
WHITE = "FFFFFF"
BORDER = "B7D7CC"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color=BORDER, size=8):
    tbl_pr = table._tbl.tblPr
    node = tbl_pr.first_child_found_in("w:tblBorders")
    if node is None:
        node = OxmlElement("w:tblBorders")
        tbl_pr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = node.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            node.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        el.set(qn("w:space"), "0")


def no_borders(table):
    tbl_pr = table._tbl.tblPr
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        node.append(el)
    tbl_pr.append(node)


def fixed_widths(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for idx, width in enumerate(widths):
        table.columns[idx].width = width
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tc_mar.find(qn(f"w:{name}"))
        if el is None:
            el = OxmlElement(f"w:{name}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def font(run, size=10, color=TEXT, bold=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def paragraph(p, before=0, after=4, line=1.05, align=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align


def keep(p, next_=False):
    p_pr = p._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepLines"))
    if next_:
        p_pr.append(OxmlElement("w:keepNext"))


def add_text(doc, text, size=10, color=TEXT, bold=False, before=0, after=5, align=None):
    p = doc.add_paragraph()
    paragraph(p, before, after, 1.08, align)
    font(p.add_run(text), size, color, bold)
    keep(p)
    return p


def add_heading(doc, title, kicker=None):
    if kicker:
        p = doc.add_paragraph()
        paragraph(p, before=3, after=1)
        font(p.add_run(kicker.upper()), 7.5, GREEN, True)
        keep(p, True)
    p = doc.add_paragraph()
    paragraph(p, before=0, after=5, line=1)
    font(p.add_run(title), 17, GREEN_DARK, True)
    keep(p, True)
    return p


def add_bullet(doc, lead, body):
    p = doc.add_paragraph(style="List Bullet")
    paragraph(p, after=3, line=1.04)
    font(p.add_run(lead), 9.6, INK, True)
    font(p.add_run(body), 9.6, TEXT)
    keep(p)
    return p


def add_callout(doc, label, body, fill=MINT):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(17.8)
    no_borders(table)
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell_margins(cell, 150, 190, 150, 190)
    p = cell.paragraphs[0]
    paragraph(p, after=2, line=1.02)
    font(p.add_run(label.upper() + "  "), 8, GREEN, True)
    font(p.add_run(body), 10.2, INK, True)
    return table


def page_break(doc):
    doc.add_page_break()


def add_page_label(doc, number, label):
    table = doc.add_table(rows=1, cols=2)
    fixed_widths(table, [Cm(1.1), Cm(16.7)])
    no_borders(table)
    c0, c1 = table.rows[0].cells
    shade(c0, GREEN)
    cell_margins(c0, 90, 90, 90, 90)
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(p, after=0)
    font(p.add_run(str(number)), 12, WHITE, True)
    cell_margins(c1, 90, 140, 90, 90)
    p = c1.paragraphs[0]
    paragraph(p, after=0)
    font(p.add_run(label.upper()), 8.5, GREEN, True)
    return table


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.35)
    sec.bottom_margin = Cm(1.35)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)
    sec.header_distance = Cm(0.55)
    sec.footer_distance = Cm(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor.from_string(TEXT)
    styles["List Bullet"].paragraph_format.left_indent = Cm(0.55)
    styles["List Bullet"].paragraph_format.first_line_indent = Cm(-0.3)

    header = sec.header
    table = header.add_table(rows=1, cols=2, width=Cm(17.8))
    table.autofit = False
    table.columns[0].width = Cm(11.7)
    table.columns[1].width = Cm(6.1)
    no_borders(table)
    for cell in table.rows[0].cells:
        cell_margins(cell, 50, 0, 60, 0)
    p = table.cell(0, 0).paragraphs[0]
    paragraph(p, after=0)
    font(p.add_run("LAND"), 9, GREEN_DARK, True)
    font(p.add_run("VERIFY"), 9, GREEN, True)
    font(p.add_run("CM"), 9, GOLD, True)
    p = table.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph(p, after=0)
    font(p.add_run("PROJECT PRESENTATION BRIEF"), 7.5, MUTED, True)

    footer = sec.footer
    p = footer.paragraphs[0]
    paragraph(p, after=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("PKFokam Institute of Excellence  •  Spring 2026  •  "), 7.2, MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)
    return doc


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = setup_document()
    core = doc.core_properties
    core.title = "Tamaffo Fabrice — Database & Project Lead Summary"
    core.subject = "LandVerifyCM capstone project presentation"
    core.author = "LandVerifyCM Team"
    core.keywords = "LandVerifyCM, Cameroon, PostgreSQL, Prisma, Database"

    # PAGE 1 — OPENING
    add_page_label(doc, 1, "Opening & project vision")
    p = doc.add_paragraph()
    paragraph(p, before=20, after=4, line=0.95)
    font(p.add_run("TAMAFFO\nFABRICE"), 28, GREEN_DARK, True)
    p = doc.add_paragraph()
    paragraph(p, after=13)
    font(p.add_run("Project Leader  •  Database Manager"), 12, GREEN, True)

    add_heading(doc, "Opening: why LandVerifyCM matters", "Suggested presentation opening")
    add_text(
        doc,
        "Good morning. Land is often a family’s lifetime investment. Yet in Cameroon, forged titles, "
        "duplicate records and unclear ownership can turn it into a dispute. LandVerifyCM makes an "
        "early check faster and more transparent.",
        11,
        INK,
        False,
        after=9,
    )
    add_callout(
        doc,
        "Our answer",
        "A centralized platform for citizens to search land records and authorized administrators to manage parcel information.",
    )

    add_heading(doc, "The project in one view", "Purpose")
    add_bullet(doc, "Public confidence — ", "search by title, parcel ID or owner; view status, location and history.")
    add_bullet(doc, "Fraud warning — ", "classify records using title uniqueness and GPS proximity.")
    add_bullet(doc, "Accountability — ", "protect record management by role and trace important changes.")
    add_bullet(doc, "Accessible design — ", "connect web and mobile interfaces to one API with English/French support.")

    add_heading(doc, "My leadership role", "Coordination")
    add_text(
        doc,
        "As project leader, I aligned three workstreams: Nkam Titcha on frontend/mobile, Kemgang Leprince "
        "on backend services, and my database work. I coordinated shared data contracts so every layer "
        "used one trusted structure.",
        9.8,
        after=0,
    )

    # PAGE 2 — DATABASE WORK
    page_break(doc)
    add_page_label(doc, 2, "Database design & implementation")
    add_heading(doc, "My responsibility: trustworthy data", "Database manager")
    add_text(
        doc,
        "I designed the PostgreSQL structure through Prisma, prepared migrations and seed data, and connected "
        "it to verification and administration. The database preserves identity, history, evidence and responsibility.",
        10.5,
        INK,
        after=7,
    )

    rows = [
        ("User", "Accounts, ADMIN/USER role, active state and secure password hash."),
        ("LandParcel", "Unique title, owner, quarter, area, GPS, land use and verification status."),
        ("OwnershipRecord", "Chronological chain: original ownership, purchase, inheritance, donation or court order."),
        ("LandDocument", "Supporting PDF/image metadata linked to its parcel."),
        ("AuditLog", "Who performed create, update, deactivation or ownership transfer, and what changed."),
    ]
    table = doc.add_table(rows=1, cols=2)
    fixed_widths(table, [Cm(4.2), Cm(13.6)])
    borders(table)
    for i, text in enumerate(("DATA MODEL", "WHAT IT CONTRIBUTES")):
        cell = table.rows[0].cells[i]
        shade(cell, GREEN_DARK)
        cell_margins(cell)
        p = cell.paragraphs[0]
        paragraph(p, after=0)
        font(p.add_run(text), 8, WHITE, True)
    for idx, (name, desc) in enumerate(rows):
        cells = table.add_row().cells
        for cell in cells:
            cell_margins(cell)
            shade(cell, MINT if idx % 2 == 0 else WHITE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[0].paragraphs[0]
        paragraph(p, after=0)
        font(p.add_run(name), 9, GREEN_DARK, True)
        p = cells[1].paragraphs[0]
        paragraph(p, after=0, line=1.0)
        font(p.add_run(desc), 8.6, TEXT)

    add_heading(doc, "How the data stays reliable", "Implemented safeguards")
    add_bullet(doc, "Integrity — ", "UUID keys, foreign keys and unique title numbers prevent ambiguity.")
    add_bullet(doc, "Traceability — ", "documents, ownership and audit events link to the correct parcel.")
    add_bullet(doc, "Performance — ", "indexes accelerate quarter, status, active-record and relation lookups.")
    add_bullet(doc, "Safe lifecycle — ", "deactivation retains records; related evidence follows controlled cascade rules.")

    add_heading(doc, "Languages and technologies—and why", "Technical choices")
    tech = doc.add_table(rows=1, cols=3)
    widths = [Cm(3.5), Cm(4.4), Cm(9.9)]
    fixed_widths(tech, widths)
    borders(tech)
    for i, text in enumerate(("CHOICE", "USED FOR", "WHY IT FITS")):
        c = tech.rows[0].cells[i]
        shade(c, GREEN)
        cell_margins(c, 90, 110, 90, 110)
        p = c.paragraphs[0]
        paragraph(p, after=0)
        font(p.add_run(text), 7.5, WHITE, True)
    tech_rows = [
        ("PostgreSQL + SQL", "Storage & migrations", "Strong constraints, joins and durable structured records."),
        ("Prisma ORM", "Schema and access", "Type-safe queries, relations and repeatable migrations."),
        ("TypeScript", "Seed and integration", "Shared types reduce API-to-database mistakes."),
    ]
    for idx, row in enumerate(tech_rows):
        cells = tech.add_row().cells
        for cell in cells:
            cell_margins(cell, 90, 110, 90, 110)
            shade(cell, WHITE if idx % 2 else MINT)
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            paragraph(p, after=0, line=1)
            font(p.add_run(text), 8.1, INK if i == 0 else TEXT, i == 0)

    # PAGE 3 — IMPACT AND CLOSE
    page_break(doc)
    add_page_label(doc, 3, "Cameroon impact & conclusion")
    add_heading(doc, "From a search to a trusted result", "End-to-end data flow")
    flow = doc.add_table(rows=1, cols=4)
    fixed_widths(flow, [Cm(4.45)] * 4)
    no_borders(flow)
    flow_items = [
        ("01", "CAPTURE", "Admin submits parcel, ownership and evidence."),
        ("02", "CHECK", "API validates values and compares title/GPS."),
        ("03", "STORE", "Prisma writes linked PostgreSQL records."),
        ("04", "TRACE", "Audit history records responsible actions."),
    ]
    for idx, (num, title, body) in enumerate(flow_items):
        c = flow.rows[0].cells[idx]
        shade(c, GREEN_DARK if idx == 0 else MINT)
        cell_margins(c, 150, 130, 150, 130)
        p = c.paragraphs[0]
        paragraph(p, after=2)
        font(p.add_run(num + "\n"), 16, GOLD, True)
        font(p.add_run(title + "\n"), 8, WHITE if idx == 0 else GREEN, True)
        font(p.add_run(body), 8.2, WHITE if idx == 0 else TEXT)

    add_heading(doc, "How this can help Cameroonian people", "Practical value")
    add_bullet(doc, "Before payment — ", "buyers can check beyond the seller’s paper documents.")
    add_bullet(doc, "Earlier fraud detection — ", "duplicate titles or overlapping GPS locations trigger warnings.")
    add_bullet(doc, "Clearer ownership — ", "families and professionals can follow transfers over time.")
    add_bullet(doc, "Better access — ", "search, maps and bilingual interfaces simplify land information.")
    add_bullet(doc, "Accountable records — ", "audit logs connect sensitive actions to an administrator.")

    add_callout(
        doc,
        "Responsible scope",
        "LandVerifyCM supports due diligence and informed decisions; formal ownership remains subject to verification by Cameroon’s competent land authorities.",
        "FFF8E8",
    )

    add_heading(doc, "Conclusion", "Suggested closing")
    add_text(
        doc,
        "To conclude, LandVerifyCM applies software engineering to a real national concern. The interfaces "
        "make verification approachable, the backend applies rules, and the database preserves evidence "
        "and history. As project leader and database manager, I connected these parts around one source of truth.",
        10.6,
        INK,
        after=7,
    )
    add_text(
        doc,
        "For Cameroonians, this means fewer blind decisions, earlier warnings and clearer information before "
        "committing hard-earned money. The platform does not replace public authorities; it provides a stronger "
        "first check and a transparent path to verification.",
        10.6,
        INK,
        after=7,
    )
    add_text(doc, "Thank you.", 13, GREEN, True, after=0)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
