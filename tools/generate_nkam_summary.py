from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables" / "team_presentations"
DOCX_PATH = OUT_DIR / "Nkam_Titcha_Frontend_Mobile_Summary.docx"

GREEN_DARK = "123D32"
GREEN = "13795B"
GREEN_BRIGHT = "20A57A"
MINT = "EAF7F1"
MINT_LIGHT = "F5FBF8"
GOLD = "D69B18"
INK = "19352E"
TEXT = "334E47"
MUTED = "62766F"
WHITE = "FFFFFF"
BORDER = "B8D8CC"


def set_run(run, size=9.5, color=TEXT, bold=False, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def format_p(p, before=0, after=0, line=1.06, align=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align


def keep(p, next_=False):
    p_pr = p._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepLines"))
    if next_:
        p_pr.append(OxmlElement("w:keepNext"))


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, top=150, start=190, bottom=150, end=190):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table, color=BORDER, size=8, inside=True):
    tbl_pr = table._tbl.tblPr
    nodes = tbl_pr.first_child_found_in("w:tblBorders")
    if nodes is None:
        nodes = OxmlElement("w:tblBorders")
        tbl_pr.append(nodes)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single" if (inside or not edge.startswith("inside")) else "nil")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        nodes.append(node)


def no_borders(table):
    borders(table, color=WHITE, size=0, inside=False)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def page_border(section):
    sect_pr = section._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "10")
        node.set(qn("w:space"), "14")
        node.set(qn("w:color"), BORDER)
        pg_borders.append(node)
    sect_pr.append(pg_borders)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    set_run(run, 7.5, MUTED, True)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_label(doc, text, number=None):
    p = doc.add_paragraph()
    format_p(p, before=7, after=4)
    keep(p, True)
    if number:
        r = p.add_run(f"{number}  ")
        set_run(r, 8.5, GOLD, True)
    r = p.add_run(text.upper())
    set_run(r, 9.5, GREEN, True)
    r.font.all_caps = True
    return p


def add_body(doc, text, size=9.2, after=5):
    p = doc.add_paragraph()
    format_p(p, after=after, line=1.08)
    keep(p)
    r = p.add_run(text)
    set_run(r, size, TEXT)
    return p


def add_bullet(container, lead, text, after=3.2, size=8.8):
    p = container.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.43)
    p.paragraph_format.first_line_indent = Cm(-0.31)
    format_p(p, after=after, line=1.06)
    keep(p)
    r = p.add_run("●  ")
    set_run(r, 6.5, GREEN_BRIGHT, True)
    r = p.add_run(lead)
    set_run(r, size, INK, True)
    r = p.add_run(text)
    set_run(r, size, TEXT)
    return p


def add_banner(doc, page_tag, title, subtitle):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.rows[0].height = Cm(3.25)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    no_borders(table)
    left, right = table.rows[0].cells
    set_cell_width(left, 13.7)
    set_cell_width(right, 4.1)
    shade(left, GREEN_DARK)
    shade(right, GREEN)
    margins(left, 260, 300, 230, 260)
    margins(right, 260, 180, 230, 180)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    p = left.paragraphs[0]
    format_p(p, after=3)
    r = p.add_run("LANDVERIFY")
    set_run(r, 9, WHITE, True)
    r = p.add_run("CM")
    set_run(r, 9, "79E2B8", True)
    p = left.add_paragraph()
    format_p(p, after=3, line=0.95)
    r = p.add_run(title)
    set_run(r, 23, WHITE, True)
    p = left.add_paragraph()
    format_p(p)
    r = p.add_run(subtitle)
    set_run(r, 9.2, "D8EEE5", False)

    p = right.paragraphs[0]
    format_p(p, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(page_tag)
    set_run(r, 8, "C5F7D6", True)
    p = right.add_paragraph()
    format_p(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("NKAM\nTITCHA")
    set_run(r, 13, WHITE, True)
    return table


def add_feature_card(doc, title, subtitle, bullets, fill=MINT_LIGHT):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    borders(table, BORDER, 8)
    cell = table.cell(0, 0)
    shade(cell, fill)
    margins(cell, 180, 230, 165, 230)
    p = cell.paragraphs[0]
    format_p(p, after=2)
    keep(p, True)
    r = p.add_run(title)
    set_run(r, 11.5, INK, True)
    p = cell.add_paragraph()
    format_p(p, after=6)
    keep(p, True)
    r = p.add_run(subtitle)
    set_run(r, 8.2, MUTED)
    for lead, text in bullets:
        add_bullet(cell, lead, text)
    return table


def add_tech_grid(doc):
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    borders(table, BORDER, 7)
    items = [
        ("React + TypeScript", "Reusable, typed components reduce UI errors and ease web/mobile maintenance."),
        ("Vite + Tailwind CSS", "Fast web builds and responsive layouts for phones, tablets and desktops."),
        ("React Router + Context", "Clear routes plus shared authentication and English/French state across the web application."),
        ("Leaflet + OpenStreetMap", "GPS parcels become understandable visual locations, with street/satellite views on the web."),
        ("Expo + React Native", "One codebase targets Android, iOS and the browser with a native-style experience."),
        ("SecureStore + Expo tools", "Protects sessions and supports document picking, downloading, sharing and embedded maps."),
    ]
    for cell, (name, why) in zip([c for row in table.rows for c in row.cells], items):
        shade(cell, MINT_LIGHT)
        margins(cell, 145, 170, 145, 170)
        p = cell.paragraphs[0]
        format_p(p, after=2)
        keep(p, True)
        r = p.add_run(name)
        set_run(r, 8.8, GREEN, True)
        p = cell.add_paragraph()
        format_p(p, line=1.02)
        keep(p)
        r = p.add_run(why)
        set_run(r, 7.9, TEXT)
    return table


def configure_doc(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.52)
    page_border(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(0)

    header = section.header
    p = header.paragraphs[0]
    format_p(p, align=WD_ALIGN_PARAGRAPH.RIGHT)
    r = p.add_run("PKFOKAM INSTITUTE OF EXCELLENCE  •  CAPSTONE 2026")
    set_run(r, 7.2, MUTED, True)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Cm(18.1))
    table.autofit = False
    no_borders(table)
    p = table.cell(0, 0).paragraphs[0]
    format_p(p)
    r = p.add_run("LANDVERIFYCM  •  FRONTEND & MOBILE BRIEF")
    set_run(r, 7.5, GREEN, True)
    add_page_number(table.cell(0, 1).paragraphs[0])


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_doc(doc)

    add_banner(
        doc,
        "TEAM MEMBER 01",
        "Frontend & Mobile Development",
        "How the citizen-facing experience turns verified records into clear, practical decisions",
    )

    add_label(doc, "Project in one view", "01")
    add_body(
        doc,
        "LandVerifyCM is a bilingual platform built to reduce land fraud and uncertainty in Cameroon. Citizens and "
        "professionals search the registry and review a parcel’s status, location and ownership history; authorized "
        "administrators maintain records. The interface connects to the Express API, where verification happens, and "
        "presents the result clearly.",
    )

    add_label(doc, "My responsibility", "02")
    add_body(
        doc,
        "As Frontend Developer, I built the interaction layer for the React website and Expo/React Native client. My "
        "work turns API data into responsive screens, protects role-specific routes, validates input and keeps the "
        "experience consistent in English and French.",
    )

    add_label(doc, "React web frontend", "03")
    add_feature_card(
        doc,
        "A guided citizen journey",
        "From first search to an informed view of a parcel",
        [
            ("Find records. ", "Users search by title number, parcel ID or owner name, or browse by quarter."),
            ("Understand results. ", "Cards and colour-coded badges show Valid, Suspicious or Duplicate status alongside loading, empty and error states."),
            ("Inspect evidence. ", "Details include area, land use, approval year, GPS coordinates, ownership history and documents."),
            ("See the location. ", "Leaflet maps place parcels on OpenStreetMap, offer street/satellite layers and copy coordinates."),
        ],
    )
    p = doc.add_paragraph()
    format_p(p, after=0)

    add_feature_card(
        doc,
        "Secure administration",
        "Tools for trusted staff to keep the public registry useful",
        [
            ("Access control. ", "Registration, login, password recovery and profile screens support users; protected routes restrict management to administrators."),
            ("Record management. ", "Admins upload, edit, deactivate and review parcels with GPS points, ownership chains and PDF/image evidence."),
            ("Usability. ", "Form validation, toast feedback, motion and responsive Tailwind layouts simplify complex tasks."),
        ],
        fill="EEF7F8",
    )

    doc.add_page_break()

    add_banner(
        doc,
        "DELIVERY NOTE",
        "Mobile Access & Public Impact",
        "One shared service, presented for practical use on phones and in a browser",
    )

    add_label(doc, "Expo / React Native client", "04")
    add_feature_card(
        doc,
        "Mobile-first access",
        "The same core service is available through Android, iOS and a browser build",
        [
            ("Navigate quickly. ", "Top and bottom navigation reach Home, Search, Browse, About and Profile; admins also receive management screens."),
            ("Verify on the move. ", "Users search, inspect status, GPS maps and ownership history, then open Google Maps."),
            ("Handle documents. ", "Authenticated users download and share evidence; admins select up to five PDF/JPG/PNG files."),
            ("Stay secure. ", "JWT sessions use SecureStore on native devices, role checks guard admin routes, and friendly errors explain failures."),
        ],
    )

    add_label(doc, "Languages and tools — and why", "05")
    add_tech_grid(doc)

    add_label(doc, "Value for Cameroon", "06")
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    borders(table, BORDER, 8)
    left, right = table.rows[0].cells
    for cell in (left, right):
        margins(cell, 170, 190, 160, 190)
    shade(left, GREEN_DARK)
    shade(right, MINT)
    p = left.paragraphs[0]
    format_p(p, after=5)
    r = p.add_run("PUBLIC BENEFIT")
    set_run(r, 9, "79E2B8", True)
    p = left.add_paragraph()
    format_p(p, line=1.05)
    r = p.add_run(
        "Phone and web access helps people check a parcel before paying, compare title information with its GPS "
        "location and notice suspicious or duplicate records earlier."
    )
    set_run(r, 8.5, WHITE)
    p = right.paragraphs[0]
    format_p(p, after=5)
    r = p.add_run("TRUST THROUGH CLARITY")
    set_run(r, 9, GREEN, True)
    p = right.add_paragraph()
    format_p(p, line=1.05)
    r = p.add_run(
        "Bilingual screens widen access; visible status, history and evidence support informed discussions with land "
        "offices and professionals. The platform supports due diligence; it does not replace official legal checks."
    )
    set_run(r, 8.5, TEXT)

    add_label(doc, "Presentation close", "07")
    p = doc.add_paragraph()
    format_p(p, after=0, line=1.08)
    r = p.add_run("My contribution makes the project’s verification work ")
    set_run(r, 9.2, TEXT)
    r = p.add_run("visible, understandable and usable")
    set_run(r, 9.2, GREEN, True)
    r = p.add_run("—especially for people accessing the service from a mobile phone.")
    set_run(r, 9.2, TEXT)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
