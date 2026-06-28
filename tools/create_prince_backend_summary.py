from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables" / "team_presentations"
DOCX_PATH = OUT_DIR / "Kemgang_Prince_Backend_Summary.docx"

GREEN_DARK = "123D32"
GREEN = "13795B"
GREEN_LIGHT = "DDF2E9"
MINT = "F1F9F5"
GOLD = "D59A16"
INK = "19352E"
TEXT = "385149"
MUTED = "6B7F78"
WHITE = "FFFFFF"
LINE = "C8DED5"
SOFT_GOLD = "FFF6DD"


def set_run(run, size=9, color=TEXT, bold=False, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=110, start=150, bottom=110, end=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    table._tbl.tblPr.append(borders)


def table_borders(table, color=LINE, size=8):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)
        node.set(qn("w:space"), "0")
        borders.append(node)
    table._tbl.tblPr.append(borders)


def keep(paragraph, next_paragraph=False):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepLines"))
    if next_paragraph:
        p_pr.append(OxmlElement("w:keepNext"))


def set_spacing(paragraph, before=0, after=0, line=1.0, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_text(parent, text, size=9, color=TEXT, bold=False, before=0, after=4, line=1.08):
    p = parent.add_paragraph()
    set_spacing(p, before, after, line)
    keep(p)
    set_run(p.add_run(text), size, color, bold)
    return p


def add_bullet(parent, lead, text, size=8.7, after=3.1):
    p = parent.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.42)
    p.paragraph_format.first_line_indent = Cm(-0.34)
    set_spacing(p, after=after, line=1.05)
    keep(p)
    set_run(p.add_run("•  "), size + 0.6, GREEN, True)
    set_run(p.add_run(lead), size, INK, True)
    set_run(p.add_run(text), size, TEXT)
    return p


def add_section_heading(parent, number, title, subtitle=None):
    p = parent.add_paragraph()
    set_spacing(p, before=2, after=2)
    keep(p, True)
    set_run(p.add_run(f"{number}  "), 8.5, GREEN, True)
    set_run(p.add_run(title.upper()), 12.5, INK, True)
    if subtitle:
        sub = parent.add_paragraph()
        set_spacing(sub, after=5, line=1.02)
        keep(sub, True)
        set_run(sub.add_run(subtitle), 8.2, MUTED)


def add_callout(parent, label, text, fill=GREEN_DARK, accent="A9E7D0"):
    table = parent.add_table(rows=1, cols=1)
    remove_borders(table)
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell_margins(cell, 90, 140, 90, 140)
    p = cell.paragraphs[0]
    set_spacing(p, line=1.02)
    set_run(p.add_run(label.upper() + "  "), 7.5, accent, True)
    set_run(p.add_run(text), 8.2, WHITE, True)
    return table


def add_page_header(doc, page_no, kicker, title, subtitle):
    band = doc.add_table(rows=1, cols=2)
    band.autofit = False
    remove_borders(band)
    left, right = band.rows[0].cells
    left.width = Cm(15.4)
    right.width = Cm(2.5)
    shade(left, GREEN_DARK)
    shade(right, GREEN)
    cell_margins(left, 170, 220, 155, 180)
    cell_margins(right, 170, 80, 155, 80)

    p = left.paragraphs[0]
    set_spacing(p, after=3)
    set_run(p.add_run(kicker.upper()), 7.5, "A9E7D0", True)
    p = left.add_paragraph()
    set_spacing(p, after=2, line=0.95)
    set_run(p.add_run(title), 22, WHITE, True)
    p = left.add_paragraph()
    set_spacing(p, line=1.0)
    set_run(p.add_run(subtitle), 8.5, "DDEDE7")

    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = right.paragraphs[0]
    set_spacing(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run(p.add_run(f"0{page_no}"), 21, WHITE, True)

    spacer = doc.add_paragraph()
    set_spacing(spacer, after=1)


def set_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_footer(section):
    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Cm(17.9))
    table.autofit = False
    remove_borders(table)
    left, right = table.rows[0].cells
    set_cell_width(left, 13.8)
    set_cell_width(right, 4.1)
    p = left.paragraphs[0]
    set_spacing(p)
    set_run(p.add_run("LANDVERIFYCM  •  PKFOKAM INSTITUTE OF EXCELLENCE"), 7.2, MUTED, True)
    p = right.paragraphs[0]
    set_spacing(p, align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_run(p.add_run("KEMGANG PRINCE  |  BACKEND"), 7.2, GREEN, True)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run(run, 7.2, MUTED, True)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.55)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    # Light green page border.
    sect_pr = section._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "10")
        el.set(qn("w:space"), "18")
        el.set(qn("w:color"), LINE)
        pg_borders.append(el)
    sect_pr.append(pg_borders)

    footer = section.footer
    p = footer.paragraphs[0]
    set_spacing(p, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run(p.add_run("BACKEND DEVELOPMENT SUMMARY   •   "), 7.2, MUTED, True)
    add_page_number_field(p)
    add_footer(section)


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    # PAGE 1 — PROJECT + ROLE
    add_page_header(
        doc,
        1,
        "LandVerifyCM • Team Presentation",
        "Backend Development Summary",
        "Kemgang Prince (Leprince)  •  Backend Developer",
    )
    add_section_heading(
        doc,
        "01",
        "Project at a glance",
        "A practical digital response to forged titles, overlapping claims, and difficult land-record checks.",
    )
    add_text(
        doc,
        "LandVerifyCM is a secure web-and-mobile land verification platform for Cameroon. Citizens can search "
        "active parcels by title number, browse by quarter, and inspect ownership, GPS, status, and documents. "
        "Administrators maintain the registry through controlled workflows. The backend bridges the React/Expo "
        "interfaces and PostgreSQL: it applies rules, protects operations, and returns consistent JSON.",
        size=9.2,
        after=7,
        line=1.12,
    )

    role_table = doc.add_table(rows=1, cols=2)
    role_table.autofit = False
    remove_borders(role_table)
    left, right = role_table.rows[0].cells
    set_cell_width(left, 8.75)
    set_cell_width(right, 8.75)
    shade(left, MINT)
    shade(right, "EEF7F8")
    for cell in (left, right):
        cell_margins(cell, 155, 175, 145, 175)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    add_section_heading(left, "02", "My concrete responsibilities")
    add_bullet(left, "REST API — ", "organized authentication, public land, and administrator routes.")
    add_bullet(left, "Business logic — ", "implemented controllers for accounts, parcels, ownership transfers, users, and audit history.")
    add_bullet(left, "Verification — ", "connected title/GPS checks to record creation and updates.")
    add_bullet(left, "Integration — ", "served stable responses consumed by the browser and mobile application.")

    add_section_heading(right, "03", "Main API surfaces")
    add_bullet(right, "Public — ", "register, login, password recovery, title search, quarter browsing, parcel details, and quarter lists.")
    add_bullet(right, "Admin — ", "list, upload, edit, deactivate, attach documents, manage ownership records, inspect users, and read audit logs.")
    add_bullet(right, "Operations — ", "startup checks, protected uploads, health monitoring, and error handling.")

    # PAGE 1 CONTINUED — ARCHITECTURE + VERIFICATION
    add_section_heading(
        doc,
        "04",
        "From request to verified record",
        "Each layer has one clear responsibility, which keeps the system easier to test, secure, and extend.",
    )

    flow = doc.add_table(rows=1, cols=5)
    flow.autofit = False
    remove_borders(flow)
    flow_items = [
        ("1", "WEB / MOBILE", "sends HTTPS request"),
        ("2", "EXPRESS ROUTE", "selects endpoint"),
        ("3", "MIDDLEWARE", "checks input and access"),
        ("4", "CONTROLLER", "applies business rules"),
        ("5", "PRISMA + SQL", "reads or writes data"),
    ]
    for cell, (number, title, detail) in zip(flow.rows[0].cells, flow_items):
        set_cell_width(cell, 3.48)
        shade(cell, MINT if int(number) % 2 else GREEN_LIGHT)
        cell_margins(cell, 115, 80, 105, 80)
        p = cell.paragraphs[0]
        set_spacing(p, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p.add_run(number), 11, GREEN, True)
        p = cell.add_paragraph()
        set_spacing(p, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p.add_run(title), 7.2, INK, True)
        p = cell.add_paragraph()
        set_spacing(p, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p.add_run(detail), 7, MUTED)

    add_text(
        doc,
        "Routes separate /api/auth, /api/lands, and /api/admin. Controllers handle requests, then Prisma sends typed "
        "queries to PostgreSQL. Pagination controls lists; selected fields protect passwords. Parcel details join "
        "documents, uploader information, and ordered ownership history.",
        size=8.8,
        before=6,
        after=7,
    )

    verify_table = doc.add_table(rows=1, cols=2)
    verify_table.autofit = False
    remove_borders(verify_table)
    rules, explanation = verify_table.rows[0].cells
    set_cell_width(rules, 8.75)
    set_cell_width(explanation, 8.75)
    shade(rules, GREEN_DARK)
    shade(explanation, SOFT_GOLD)
    for cell in (rules, explanation):
        cell_margins(cell, 160, 180, 150, 180)

    p = rules.paragraphs[0]
    set_spacing(p, after=6)
    set_run(p.add_run("VERIFICATION DECISION RULES"), 11, WHITE, True)
    for line in (
        "1  Matching title number → DUPLICATE",
        "2  GPS distance under 10 m → DUPLICATE",
        "3  GPS distance under 50 m → SUSPICIOUS",
        "4  No conflict → VALID",
    ):
        p = rules.add_paragraph()
        set_spacing(p, after=4, line=1.03)
        set_run(p.add_run(line), 8.5, WHITE, True)

    add_section_heading(explanation, "05", "Why the algorithm matters")
    add_text(
        explanation,
        "The Haversine formula converts GPS pairs into distance using Earth’s radius (6,371 km). Checks compare active "
        "parcels and exclude the current parcel during updates. A duplicate upload returns HTTP 409; a nearby record "
        "is stored as Suspicious for manual review.",
        size=8.5,
        after=0,
        line=1.08,
    )

    add_callout(
        doc,
        "Design value",
        "One verification service means the browser and phone receive the same evidence-based decision.",
        fill="245B68",
        accent="BDEEF1",
    )

    doc.add_page_break()

    # PAGE 2 — SECURITY + STACK + IMPACT
    add_page_header(
        doc,
        2,
        "Trust, technology, and impact",
        "Built for Safe Public Use",
        "Security controls protect the registry while the technology keeps development practical.",
    )
    top = doc.add_table(rows=1, cols=2)
    top.autofit = False
    remove_borders(top)
    security, stack = top.rows[0].cells
    set_cell_width(security, 8.75)
    set_cell_width(stack, 8.75)
    shade(security, MINT)
    shade(stack, "EEF7F8")
    for cell in (security, stack):
        cell_margins(cell, 145, 170, 135, 170)

    add_section_heading(security, "06", "Security and accountability")
    add_bullet(security, "Identity — ", "bcrypt uses 12 rounds; JWT carries user ID and role.")
    add_bullet(security, "Access — ", "Bearer authentication and ADMIN authorization protect management and files.")
    add_bullet(security, "Abuse controls — ", "Helmet, CORS, body limits, and rate limits reduce attacks.")
    add_bullet(security, "Evidence — ", "important changes receive timestamped audit records.")
    add_bullet(security, "Uploads — ", "Multer permits five PDF/JPG/PNG files, maximum 10 MB each.")

    add_section_heading(stack, "07", "Technologies and why")
    add_bullet(stack, "TypeScript — ", "catches data-shape mistakes before deployment.")
    add_bullet(stack, "Node.js + Express — ", "delivers a lightweight REST service.")
    add_bullet(stack, "PostgreSQL + Prisma — ", "provides relational storage, constraints, migrations, and typed queries.")
    add_bullet(stack, "Zod — ", "validates coordinates, years, areas, and ownership data.")
    add_bullet(stack, "JWT, bcrypt, Multer — ", "support sessions, safe passwords, and controlled uploads.")

    add_section_heading(
        doc,
        "08",
        "How this helps people in Cameroon",
        "The backend turns scattered land information into checks citizens and administrators can act on.",
    )
    impact = doc.add_table(rows=1, cols=3)
    impact.autofit = False
    remove_borders(impact)
    impact_items = [
        ("BUYERS", "Check a title and location before paying, reducing exposure to forged or overlapping claims."),
        ("FAMILIES & PROFESSIONALS", "Follow ownership history and supporting evidence when investigating a parcel."),
        ("ADMINISTRATORS", "Maintain records with controlled changes, traceable actions, and clear manual-review flags."),
    ]
    for i, (cell, (label, text)) in enumerate(zip(impact.rows[0].cells, impact_items)):
        set_cell_width(cell, 5.82)
        shade(cell, (MINT, SOFT_GOLD, GREEN_LIGHT)[i])
        cell_margins(cell, 125, 125, 115, 125)
        p = cell.paragraphs[0]
        set_spacing(p, after=4)
        set_run(p.add_run(label), 7.7, GREEN_DARK, True)
        p = cell.add_paragraph()
        set_spacing(p, line=1.05)
        set_run(p.add_run(text), 8, TEXT)

    add_callout(
        doc,
        "Prince’s contribution",
        "A consistent and secure backend that converts land records into useful verification results for both web and mobile users.",
    )

    doc.core_properties.title = "Kemgang Prince — LandVerifyCM Backend Development Summary"
    doc.core_properties.subject = "Backend role, architecture, verification, security, technology, and public impact"
    doc.core_properties.author = "LandVerifyCM Project Team"
    doc.core_properties.keywords = "LandVerifyCM, Cameroon, backend, Express, TypeScript, Prisma, PostgreSQL"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
